"""Generate phase-wise .docx, xlsx, pptx, diagrams, README extras."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent.parent
SHOT = ROOT / "docs" / "screenshots"
PH = ROOT / "Phase_Wise_Submission"
DIAG = PH / "03_Project_Design" / "diagrams"


def add_title(doc: Document, title: str, subtitle: str = ""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(19, 78, 74)
    if subtitle:
        s = doc.add_paragraph(subtitle)
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("PocketSmart AI — Your Smart Budget and Recommendation Assistant")
    doc.add_paragraph("Submitted by: [YOUR NAME]    Mentor: [MENTOR NAME]    Date: 24 September 2026")
    doc.add_page_break()
    doc.add_heading("Table of contents", level=1)
    doc.add_paragraph("(Use Word: References → Table of Contents after opening this file.)")


def footer(doc: Document):
    for section in doc.sections:
        section.footer.paragraphs[0].text = "PocketSmart AI  |  Page "
        section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def fig(doc, rel, caption):
    path = SHOT / rel
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.2))
    p = doc.add_paragraph(caption)
    p.italic = True


def save(doc, path: Path):
    footer(doc)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def phase1():
    d = Document()
    add_title(d, "Phase 01 — Brainstorming and Ideation")
    d.add_heading("1. Problem statement", 1)
    d.add_paragraph(
        "People who furnish a home, host a party, or buy jewelry for an occasion must compare prices across Amazon, Flipkart, IKEA, Swiggy, Zomato and OYO. Spreadsheets overflow and totals quietly cross the budget. PocketSmart AI turns a rupee budget plus simple preferences into a capped shopping list. Catalogs in this academic build are mock JSON labelled “sample data”; they are not live scrapes."
    )
    d.add_heading("2. Target users and personas", 1)
    table = d.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    hdr = ["Persona", "Role", "Need", "Pain"]
    for i, h in enumerate(hdr):
        table.rows[0].cells[i].text = h
    rows = [
        ["Riya Sharma", "New renter, Chennai", "Furnish 2BHK under ₹80,000", "Too many SKUs, no room-wise split"],
        ["Arjun Patel", "Event host, Pune", "Birthday for 40 guests, ₹50,000", "Catering + décor + stay overshoots"],
        ["Nisha Reddy", "Wedding guest, Hyderabad", "Jewelry under ₹15,000 matching outfit", "Guessing colours from a photo"],
    ]
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r):
            table.rows[i].cells[j].text = v
    d.add_heading("3. Five candidate ideas", 1)
    t = d.add_table(rows=6, cols=3)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Idea"
    t.rows[0].cells[1].text = "Pros"
    t.rows[0].cells[2].text = "Cons"
    ideas = [
        ("Generic expense tracker", "Simple", "No product suggestions"),
        ("Single-store IKEA bot", "Accurate SKUs", "Not cross-platform"),
        ("Scraped live marketplace", "Real prices", "Against this spec; legal risk"),
        ("Spreadsheet template", "Offline", "No AI, no images"),
        ("PocketSmart AI (chosen)", "Three planners, budget cap, demo mode", "Mock catalogs only"),
    ]
    for i, row in enumerate(ideas, 1):
        for j, v in enumerate(row):
            t.rows[i].cells[j].text = v
    d.add_paragraph("PocketSmart AI was chosen because the spec requires Flask, Gemini multimodal jewelry photos, and three planners with a hard budget rule.")
    d.add_heading("4. Proposed solution and unique value", 1)
    d.add_paragraph(
        "A Flask app on port 5000 with HTML/CSS/JS, session history, CORS, and /health. Gemini (google-genai, model from GEMINI_MODEL) is used when GEMINI_API_KEY is set; otherwise deterministic catalog pickers run (Demo mode badge). After any AI JSON, enforce_budget() drops items until the total is ≤ budget."
    )
    d.add_heading("5. Features", 1)
    d.add_paragraph("Must-have: Home / Party / Jewelry POST routes, validation, budget cap, demo fallback, sample-data labels, screenshots, pytest.")
    d.add_paragraph("Nice-to-have: Register/login session, testimonials, history page, outfit image analysis.")
    d.add_heading("6. Why Gemini, and risks", 1)
    d.add_paragraph(
        "Gemini fits because jewelry needs optional image understanding and planners need structured JSON. Risks: wrong SKUs, API cost, privacy of outfit photos. Mitigations: mock catalogs only, demo mode, no stored uploads, fallback lists, visible sample-data tags."
    )
    fig(d, "01_home.png", "Figure 1. Landing page used in ideation walkthroughs.")
    save(d, PH / "01_Brainstorming_Ideation" / "Phase_01_Brainstorming.docx")


def phase2():
    d = Document()
    add_title(d, "Phase 02 — Requirement Analysis")
    d.add_heading("1. Functional requirements", 1)
    t = d.add_table(rows=11, cols=3)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "ID"
    t.rows[0].cells[1].text = "Description"
    t.rows[0].cells[2].text = "Priority"
    frs = [
        ("FR-01", "POST /generate-home with budget, rooms, quantities", "Must"),
        ("FR-02", "POST /generate-party with budget, guests, event, venue", "Must"),
        ("FR-03", "POST /generate-jewelry with budget, occasion, style, optional image", "Must"),
        ("FR-04", "Total of items never exceeds budget", "Must"),
        ("FR-05", "Demo mode when GEMINI_API_KEY empty", "Must"),
        ("FR-06", "Fallback message if Gemini fails", "Must"),
        ("FR-07", "GET /health reports mode", "Must"),
        ("FR-08", "Reject zero/negative/too-small budgets", "Must"),
        ("FR-09", "Reject non-image uploads", "Must"),
        ("FR-10", "Session history of plans", "Should"),
    ]
    for i, r in enumerate(frs, 1):
        for j, v in enumerate(r):
            t.rows[i].cells[j].text = v
    d.add_heading("2. Non-functional requirements", 1)
    d.add_paragraph("Performance: demo recommendations under 2 s. Usability: 1366×768 desktop and 390×844 mobile. Security: no secrets in repo, empty .env.example. Reliability: 22/22 pytest passed (83% coverage).")
    d.add_heading("3. User stories", 1)
    stories = [
        "US-01 As Riya I enter ₹80,000 and rooms so I receive a home list under budget.",
        "US-02 As Arjun I enter 40 guests and birthday so I get catering/décor/entertainment split.",
        "US-03 As Nisha I upload an outfit PNG so jewelry matches colours.",
        "US-04 As a tester I send budget 0 and see HTTP 400.",
        "US-05 As a tester I upload notes.txt and see invalid image.",
        "US-06 As a student without an API key I still use Demo mode.",
        "US-07 As a user I open /history and see my last plans.",
        "US-08 As ops I GET /health and see mode=demo or live.",
    ]
    for s in stories:
        d.add_paragraph(s, style="List Number")
        d.add_paragraph("Acceptance: matching pytest and UI screenshot in docs/screenshots/.")
    d.add_heading("4. Technical requirements", 1)
    d.add_paragraph("Python 3.13, Flask 3.0.3, Flask-CORS 4.0.1, google-genai, Pillow 10.4, pytest 8.3. Hardware: any laptop. API key: GEMINI_API_KEY (optional). Port 5000.")
    d.add_heading("5. Constraints, assumptions, dependencies", 1)
    d.add_paragraph("Constraint: no live scraping. Assumption: INR. Dependency: mock JSON catalogs. Gemini model ID only in pocketsmart/config.py via env.")
    d.add_heading("6. Scope", 1)
    d.add_paragraph("In scope: three planners, demo/live Gemini, tests, docs, videos. Out of scope: payments, real Amazon APIs, IBM Watsonx, FastAPI (spec architecture is Flask).")
    fig(d, "05_home_planner_form.png", "Figure 1. Home planner form implementing FR-01.")
    save(d, PH / "02_Requirement_Analysis" / "Phase_02_Requirements.docx")


def diagrams():
    DIAG.mkdir(parents=True, exist_ok=True)

    def box(draw, xy, text, fill="#134e4a"):
        draw.rounded_rectangle(xy, 12, fill=fill)
        draw.text((xy[0] + 12, xy[1] + 18), text, fill="white")

    def canvas(name, title, lines):
        im = Image.new("RGB", (1100, 620), "#f6f1e8")
        dr = ImageDraw.Draw(im)
        dr.text((24, 16), title, fill="#134e4a")
        y = 70
        for line in lines:
            dr.rounded_rectangle((40, y, 1060, y + 48), 8, fill="#fffdf8", outline="#c4a35a")
            dr.text((56, y + 14), line, fill="#1c1917")
            y += 60
        im.save(DIAG / name)

    canvas(
        "01_architecture.png",
        "Figure D1 — System architecture (matches code)",
        [
            "Browser HTML/CSS/JS  →  Flask 0.0.0.0:5000 (pages + api blueprints)",
            "routes/api.py  POST /generate-home|party|jewelry  GET /health",
            "services/recommend.py  +  services/ai_client.py (google-genai or demo)",
            "pocketsmart/config.py  GEMINI_MODEL / DEMO_MODE",
            "data/*_catalog.json  mock Amazon, Flipkart, IKEA, Swiggy, Zomato, OYO",
            "Session cookie stores history[]; no database server",
        ],
    )
    canvas(
        "02_workflow.png",
        "Figure D2 — User workflow",
        [
            "Open planner form → validate budget/rooms/guests/image",
            "Call recommend_* → Gemini JSON or demo picker",
            "enforce_budget() until sum(qty*price) ≤ budget",
            "Label sample data → JSON → JS cards → session history",
        ],
    )
    canvas(
        "03_dfd.png",
        "Figure D3 — DFD level 0 / 1",
        [
            "L0: User → PocketSmart AI → Recommendations (INR, sample catalogs)",
            "L1 Home: rooms+qty → Home engine → IKEA/Amazon/Flipkart items",
            "L1 Party: guests+event → split catering/décor/entertainment/stay",
            "L1 Jewelry: occasion+style+image → Amazon/Flipkart jewelry",
        ],
    )
    canvas(
        "04_usecase.png",
        "Figure D4 — Use cases",
        [
            "Actor Student: Register, Login, Plan home, Plan party, Plan jewelry",
            "Actor Tester: Hit /health, submit invalid budget, upload .txt",
            "System: Cap budget, fallback, store session history",
        ],
    )
    canvas(
        "05_ai_flow.png",
        "Figure D5 — AI generation flow",
        [
            "Build prompt (HOME_PROMPTS / PARTY_PROMPT / JEWEL_PROMPT)",
            "If key present: google.genai Client generate_content(GEMINI_MODEL)",
            "parse_json_payload → hydrate catalog ids → enforce_budget",
            "If invalid JSON or <2 items → fallback_* generators",
        ],
    )
    canvas(
        "06_fallback.png",
        "Figure D6 — Fallback flow",
        [
            "DEMO_MODE or API exception or non-JSON → source=demo|fallback",
            "fallback_home / fallback_party / fallback_jewelry greedy cheapest",
            "UI shows yellow banner FALLBACK_MESSAGE + sample-data tags",
            "HTTP still 200 so screenshots and demos never crash",
        ],
    )


def phase3():
    diagrams()
    d = Document()
    add_title(d, "Phase 03 — Project Design")
    d.add_heading("1. Diagrams", 1)
    for name, expl in [
        ("01_architecture.png", "Flask process owns templates and JSON APIs; Gemini is optional."),
        ("02_workflow.png", "Every planner follows validate → generate → cap → display."),
        ("03_dfd.png", "No external marketplace calls; only local JSON."),
        ("04_usecase.png", "Three planners plus auth and testing actors."),
        ("05_ai_flow.png", "Prompts ask for JSON with catalog ids only."),
        ("06_fallback.png", "Demo mode is the default in this repository."),
    ]:
        pth = DIAG / name
        d.add_picture(str(pth), width=Inches(6.2))
        d.add_paragraph(expl)
    d.add_heading("2. Data model", 1)
    d.add_paragraph("No SQL database. Session JSON: user {name,email}, history list of {planner,budget,total,summary,item_count}. Catalog item: id, name, category, price (INR), platform, url.")
    d.add_heading("3. API table", 1)
    t = d.add_table(rows=8, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(["Method", "Route", "Input", "Output"]):
        t.rows[0].cells[i].text = h
    rows = [
        ("GET", "/health", "—", "mode demo|live"),
        ("POST", "/generate-home", "budget, rooms, quantities", "items, total"),
        ("POST", "/generate-party", "budget, guests, event, venue", "items, split"),
        ("POST", "/generate-jewelry", "budget, occasion, style, file", "items, outfit_analysis"),
        ("POST", "/auth/login", "email", "user"),
        ("GET", "/api/history", "cookie", "history[]"),
        ("GET", "/home-planner", "—", "HTML"),
    ]
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r):
            t.rows[i].cells[j].text = v
    d.add_heading("4. Prompt design", 1)
    d.add_paragraph("HOME_PROMPTS, PARTY_PROMPT and JEWEL_PROMPT in services/recommend.py ask for JSON with catalog ids only and “Never exceed budget.”")
    fig(d, "06_home_results.png", "Figure 7. Home results UI corresponding to the API.")
    save(d, PH / "03_Project_Design" / "Phase_03_Design.docx")


def phase4():
    wb = Workbook()
    ws = wb.active
    ws.title = "Gantt"
    headers = ["Task", "Owner", "Start", "End", "Days", "Status", "W1", "W2", "W3", "W4"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h).font = Font(bold=True)
        ws.cell(1, i).fill = PatternFill("solid", fgColor="134E4A")
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
    tasks = [
        ("Gemini setup & config.py", "Dev", "2026-09-01", "2026-09-03", 3, "Done", 1, 0, 0, 0),
        ("Flask routes & sessions", "Dev", "2026-09-04", "2026-09-07", 4, "Done", 1, 0, 0, 0),
        ("Catalogs & recommenders", "Dev", "2026-09-08", "2026-09-10", 3, "Done", 0, 1, 0, 0),
        ("HTML/CSS/JS UI", "Dev", "2026-09-11", "2026-09-14", 4, "Done", 0, 1, 0, 0),
        ("pytest 22 cases", "QA", "2026-09-15", "2026-09-17", 3, "Done", 0, 0, 1, 0),
        ("Screenshots & videos", "QA", "2026-09-18", "2026-09-20", 3, "Done", 0, 0, 1, 0),
        ("Phase documents", "Writer", "2026-09-21", "2026-09-24", 4, "Done", 0, 0, 0, 1),
        ("Zip & viva pack", "Writer", "2026-09-24", "2026-09-24", 1, "Done", 0, 0, 0, 1),
    ]
    for r, row in enumerate(tasks, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    chart = BarChart()
    chart.type = "bar"
    chart.title = "Week load"
    data = Reference(ws, min_col=7, min_row=1, max_col=10, max_row=9)
    cats = Reference(ws, min_col=1, min_row=2, max_row=9)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "A12")
    gantt_path = PH / "04_Project_Planning" / "Gantt_Chart.xlsx"
    wb.save(gantt_path)

    d = Document()
    add_title(d, "Phase 04 — Project Planning")
    d.add_heading("1. Work breakdown (spec workflow)", 1)
    d.add_paragraph("M1 Gemini setup. M2 Core Flask routes. M3 Modular services. M4 UI. M5 Testing. Activities 1.1–5.4 from the spec were mapped onto the Flask+Gemini architecture (not Watsonx/FastAPI).")
    d.add_heading("2. Four-week timeline", 1)
    d.add_paragraph("Week 1: config, catalogs, app factory. Week 2: recommenders and UI. Week 3: pytest, Playwright shots. Week 4: phase docs, videos, zip.")
    d.add_heading("3. Gantt", 1)
    d.add_paragraph("See Gantt_Chart.xlsx (tasks, owner, dates, week bars). All eight tasks are Done as of 24 Sep 2026.")
    d.add_heading("4. Risk register", 1)
    t = d.add_table(rows=5, cols=4)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Risk"
    t.rows[0].cells[1].text = "P"
    t.rows[0].cells[2].text = "I"
    t.rows[0].cells[3].text = "Mitigation"
    risks = [
        ("No Gemini key", "H", "M", "Demo mode"),
        ("AI over-budget JSON", "M", "H", "enforce_budget"),
        ("Playwright libs missing", "M", "M", "Document LD_LIBRARY_PATH"),
        ("Huge uploads", "L", "M", "MAX_UPLOAD_MB + 413"),
    ]
    for i, r in enumerate(risks, 1):
        for j, v in enumerate(r):
            t.rows[i].cells[j].text = v
    d.add_heading("5. Tools", 1)
    d.add_paragraph("VS Code / this workspace, Git, pytest, Playwright Chromium, ffmpeg (imageio-ffmpeg), python-docx, openpyxl, python-pptx.")
    save(d, PH / "04_Project_Planning" / "Phase_04_Planning.docx")


def phase5():
    d = Document()
    add_title(d, "Phase 05 — Project Development")
    d.add_heading("1. Environment", 1)
    d.add_paragraph("python3 -m venv .venv && pip install -r requirements.txt")
    d.add_paragraph("Copy .env.example to .env (leave GEMINI_API_KEY empty for demo).")
    d.add_paragraph("PYTHONPATH=. python3 -m pocketsmart.app   # port 5000")
    d.add_heading("2. Structure", 1)
    d.add_paragraph("pocketsmart/app.py — Flask factory. config.py — model + demo flag. routes/api.py — REST. routes/pages.py — HTML. services/ai_client.py — Gemini. services/recommend.py — budget engine. models/schemas.py — validation. data/*.json — catalogs. templates/ + static/.")
    d.add_heading("3. AI integration", 1)
    d.add_paragraph("google-genai Client. Model GEMINI_MODEL default gemini-2.0-flash. DEMO_MODE if key empty. parse_json_payload strips fences. Fallback lists from catalogs.")
    d.add_heading("4. Frontend", 1)
    fig(d, "06_home_results.png", "Figure 1. Home results.")
    fig(d, "08_party_results.png", "Figure 2. Party results.")
    fig(d, "10_jewelry_results.png", "Figure 3. Jewelry with outfit.")
    d.add_heading("5. Challenges (honest)", 1)
    d.add_paragraph("1. Spec mixes FastAPI/Watsonx with Flask/Gemini — implemented Flask as architecture section requires.")
    d.add_paragraph("2. Gemini 1.5 names outdated — env-driven GEMINI_MODEL.")
    d.add_paragraph("3. /history collided with HTML page — JSON moved to /api/history.")
    d.add_paragraph("4. Playwright needed extra .so libraries in this sandbox (libnspr4, libnss3, atk).")
    d.add_paragraph("5. System ffmpeg missing — used imageio-ffmpeg static binary to write MP4.")
    save(d, PH / "05_Project_Development" / "Phase_05_Development.docx")


def phase6():
    wb = Workbook()
    ws = wb.active
    ws.title = "Test_Cases"
    headers = ["ID", "Module", "Description", "Input", "Expected", "Actual", "Status", "Screenshot"]
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="134E4A")
    cases = [
        ("TC-01", "Health", "GET /health", "none", "200 mode", "200 demo", "Pass", "—"),
        ("TC-02", "Home", "₹80k three rooms", "80000", "total≤budget", "pass pytest", "Pass", "06_home_results.png"),
        ("TC-03", "Party", "birthday 40 guests", "50000", "total≤budget", "pass", "Pass", "08_party_results.png"),
        ("TC-04", "Jewelry", "wedding guest 15k", "15000", "total≤budget", "pass", "Pass", "10_jewelry_results.png"),
        ("TC-05", "Val", "budget 0", "0", "400", "400", "Pass", "13_edge_budget_too_low.png"),
        ("TC-06", "Val", "negative budget", "-10", "400", "400", "Pass", "—"),
        ("TC-07", "Val", "budget 100 home", "100", "budget_too_small", "400", "Pass", "13_edge_budget_too_low.png"),
        ("TC-08", "Val", "no rooms", "[]", "400", "400", "Pass", "—"),
        ("TC-09", "Val", "bad event type", "alien", "400", "400", "Pass", "—"),
        ("TC-10", "Val", "guests 0", "0", "400", "400", "Pass", "—"),
        ("TC-11", "Image", "notes.txt", "txt", "400 invalid_image", "400", "Pass", "14_edge_invalid_image.png"),
        ("TC-12", "Image", "PNG outfit", "png", "200 + analysis", "200", "Pass", "10_jewelry_results.png"),
        ("TC-13", "Session", "history after party", "cookie", "planner=party", "pass", "Pass", "11_history.png"),
        ("TC-14", "Auth", "login", "a@b.com", "200", "200", "Pass", "04_login.png"),
        ("TC-15", "AI", "generate_structured None", "patch", "fallback items", "pass", "Pass", "15_edge_fallback.png"),
        ("TC-16", "AI", "bad ids", "live fake", "fallback", "pass", "Pass", "—"),
        ("TC-17", "Engine", "enforce_budget 250", "900+200+50", "≤250", "pass", "Pass", "—"),
        ("TC-18", "Engine", "fallback_home 80k", "qty", "≤80000", "pass", "Pass", "—"),
        ("TC-19", "Engine", "fallback_party hotel", "40", "has catering", "pass", "Pass", "—"),
        ("TC-20", "Auth", "register no name", "{}", "400", "400", "Pass", "—"),
        ("TC-21", "Pages", "planner HTML", "GET", "200", "200", "Pass", "01_home.png"),
        ("TC-22", "UI", "Demo badge", "GET /", "Demo mode", "visible", "Pass", "01_home.png"),
    ]
    for r, row in enumerate(cases, 2):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    ws2 = wb.create_sheet("Traceability")
    ws2.append(["Requirement", "Tests"])
    mapping = [
        ("FR-01", "TC-02,TC-21"),
        ("FR-02", "TC-03"),
        ("FR-03", "TC-04,TC-12"),
        ("FR-04", "TC-02,TC-03,TC-04,TC-17,TC-18"),
        ("FR-05", "TC-01,TC-22"),
        ("FR-06", "TC-15,TC-16"),
        ("FR-07", "TC-01"),
        ("FR-08", "TC-05,TC-06,TC-07"),
        ("FR-09", "TC-11"),
        ("FR-10", "TC-13"),
    ]
    for m in mapping:
        ws2.append(list(m))
    wb.save(PH / "06_Project_Testing" / "Test_Cases.xlsx")
    wb2 = Workbook()
    w = wb2.active
    w.title = "Traceability_Matrix"
    w.append(["FR", "Test IDs"])
    for m in mapping:
        w.append(list(m))
    wb2.save(PH / "06_Project_Testing" / "Traceability_Matrix.xlsx")

    d = Document()
    add_title(d, "Phase 06 — Project Testing")
    d.add_heading("1. Strategy", 1)
    d.add_paragraph("Unit (enforce_budget), API (Flask test client), UI (Playwright), edge (budget/image), AI-failure (mock generate_structured).")
    d.add_heading("2. Summary (real pytest)", 1)
    d.add_paragraph("22 collected, 22 passed, 0 failed. Coverage 83% (pocketsmart). Evidence: docs/evidence/pytest.txt and pytest_cov.txt.")
    d.add_heading("3. Bugs found and fixed", 1)
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Bug"
    t.rows[0].cells[1].text = "Fix"
    t.rows[0].cells[2].text = "Status"
    t.rows[1].cells[0].text = "GET /history returned HTML so JSON tests failed"
    t.rows[1].cells[1].text = "Moved API to /api/history"
    t.rows[1].cells[2].text = "Fixed"
    t.rows[2].cells[0].text = "Invalid image test sent raw bytes"
    t.rows[2].cells[1].text = "BytesIO + filename notes.txt"
    t.rows[2].cells[2].text = "Fixed"
    t.rows[3].cells[0].text = "Outfit analysis missing without image_bytes"
    t.rows[3].cells[1].text = "Multipart branch + demo analysis text"
    t.rows[3].cells[2].text = "Fixed"
    fig(d, "13_edge_budget_too_low.png", "Figure 1. Budget too low.")
    fig(d, "14_edge_invalid_image.png", "Figure 2. Invalid image.")
    fig(d, "15_edge_fallback.png", "Figure 3. Fallback list.")
    save(d, PH / "06_Project_Testing" / "Phase_06_Testing.docx")


def phase7():
    d = Document()
    add_title(d, "Final Project Report", "PocketSmart AI")
    d.add_heading("Certificate (placeholder)", 1)
    d.add_paragraph("This is to certify that [YOUR NAME] has completed the academic project PocketSmart AI under the guidance of [MENTOR NAME].")
    d.add_heading("Acknowledgement", 1)
    d.add_paragraph("Thanks to faculty and to the open-source Flask, pytest and Playwright communities.")
    d.add_heading("Abstract", 1)
    d.add_paragraph(
        "PocketSmart AI is a Flask web application that plans home interiors, parties and jewelry inside a stated INR budget. Google Gemini is optional; empty GEMINI_API_KEY enables Demo mode with deterministic catalog pickers. Twenty-two pytest cases passed with 83% coverage. Mock catalogs are labelled sample data."
    )
    d.add_heading("1 Introduction", 1)
    d.add_paragraph("Cross-platform shopping for life events is noisy. This project gives three focused planners and a hard budget invariant.")
    d.add_heading("2 Literature and technology review", 1)
    d.add_paragraph("Flask for routing and sessions; Gemini multimodal for outfit photos; pytest for API contracts. IBM Watsonx and FastAPI appear in an older prerequisite list but were not used, following the architecture section of the spec.")
    d.add_heading("3 Problem and objectives", 1)
    d.add_paragraph("Objectives: (1) budget-capped lists (2) three domains (3) demo/live AI (4) tests and phase documents.")
    d.add_heading("4 System analysis", 1)
    d.add_paragraph("See Phase 02 FR-01–FR-10. Users are renters, hosts and wedding guests in India.")
    d.add_heading("5 System design", 1)
    d.add_paragraph("Modular package pocketsmart with config-driven model IDs. Diagrams live in Phase_Wise_Submission/03_Project_Design/diagrams/.")
    d.add_heading("6 Implementation", 1)
    d.add_paragraph("create_app registers pages_bp and api_bp, enables CORS and sessions. recommend.py hydrates AI ids against catalogs then enforce_budget.")
    d.add_heading("7 Testing and results", 1)
    d.add_paragraph("pytest: 22 passed. Playwright: 18 screenshots including 3 mobile. Demo video 1m35s; testing video 1m02s (shorter than the 3–4 / 2–3 minute brief because scene timers were conservative; scripts can be re-run with longer waits).")
    fig(d, "01_home.png", "Figure 1. Home.")
    fig(d, "06_home_results.png", "Figure 2. Home plan ₹80,000.")
    fig(d, "08_party_results.png", "Figure 3. Party ₹50,000 / 40 guests.")
    fig(d, "10_jewelry_results.png", "Figure 4. Jewelry ₹15,000.")
    fig(d, "16_mobile_home.png", "Figure 5. Mobile 390×844.")
    d.add_heading("8 Limitations", 1)
    d.add_paragraph("Mock prices; no payments; demo AI; videos under the requested duration; Chrome needed extra system libraries in the build sandbox.")
    d.add_heading("9 Future enhancements", 1)
    d.add_paragraph("Live price APIs with user consent, PDF export, multi-user database, GST-aware totals.")
    d.add_heading("10 Conclusion", 1)
    d.add_paragraph("PocketSmart AI shows that a small Flask service plus optional Gemini can keep lifestyle spending inside a budget when catalogs are controlled and totals are validated in code.")
    d.add_heading("References", 1)
    d.add_paragraph("Flask documentation; Google Gen AI Python SDK; pytest; Playwright.")
    d.add_heading("Appendix A — Setup", 1)
    d.add_paragraph("pip install -r requirements.txt && PYTHONPATH=. python3 -m pocketsmart.app")
    d.add_heading("Appendix B — API", 1)
    d.add_paragraph("See Phase 03 route table.")
    save(d, PH / "07_Project_Documentation" / "Final_Project_Report.docx")


def phase8():
    d = Document()
    add_title(d, "Demo Guide")
    d.add_heading("Pre-demo checklist", 1)
    d.add_paragraph("Port 5000 free. GEMINI_API_KEY empty (Demo mode badge visible). Sample outfit at pocketsmart/static/img/sample_outfit.png.")
    d.add_heading("5-minute script", 1)
    d.add_paragraph("0:00 Landing. 0:40 Home ₹80,000 living/kitchen/bedroom. 1:40 Party birthday 40 guests ₹50,000. 2:40 Jewelry wedding guest ₹15,000 + outfit. 3:40 History. 4:10 Edge budget 100. 4:40 Thank you.")
    d.add_paragraph("If API fails: yellow fallback banner is expected in demo; continue.")
    d.add_paragraph("Videos: media/Demo_Video.mp4 and media/Testing_Video.mp4")
    save(d, PH / "08_Project_Demonstration" / "Demo_Guide.docx")

    prs = Presentation()
    titles = [
        ("PocketSmart AI", "Smart budget & recommendation assistant"),
        ("Problem", "Home, party and jewelry spending spreads across platforms"),
        ("Solution", "Flask planners + optional Gemini + hard budget cap"),
        ("Architecture", "Browser → Flask:5000 → recommend.py → catalogs / Gemini"),
        ("Tech stack", "Flask, Flask-CORS, google-genai, Pillow, pytest, Playwright"),
        ("Features", "Three planners, demo badge, sample-data labels, /health"),
        ("Live demo shots", "See screenshots 06, 08, 10"),
        ("Testing", "22/22 pytest, 83% coverage"),
        ("Challenges", "Spec FastAPI vs Flask; Playwright libs; ffmpeg"),
        ("Future work", "Live prices, PDF export, accounts DB"),
        ("Conclusion", "Budget invariant in code, not only in the prompt"),
        ("Thank you", "Questions?  media/Demo_Video.mp4"),
    ]
    for t, s in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = t
        slide.placeholders[1].text = s
    # add a picture slide if screenshot exists
    if (SHOT / "06_home_results.png").exists():
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = "Home planner results"
        sl.shapes.add_picture(str(SHOT / "06_home_results.png"), PInches(0.5), PInches(1.4), width=PInches(9))
    prs.save(str(PH / "08_Project_Demonstration" / "Final_Presentation.pptx"))

    v = Document()
    add_title(v, "Viva Questions")
    qa = [
        ("What is PocketSmart AI?", "A Flask app that recommends home, party and jewelry items under an INR budget using mock catalogs and optional Gemini."),
        ("Why Flask not FastAPI?", "The architecture section specifies Flask; we followed that over the outdated prerequisites."),
        ("Port?", "5000."),
        ("How is the model selected?", "GEMINI_MODEL in .env, read by pocketsmart/config.py."),
        ("What is Demo mode?", "Empty or placeholder GEMINI_API_KEY; deterministic fallback generators; yellow badge."),
        ("How do you cap budget?", "enforce_budget sorts by line total and keeps items while sum ≤ budget."),
        ("What if Gemini returns junk JSON?", "parse_json_payload fails → fallback_*."),
        ("Image upload rules?", "png/jpg/webp, size MAX_UPLOAD_MB, else 400 invalid_image."),
        ("Are Amazon results real?", "No. JSON mocks labelled sample data."),
        ("Currency?", "INR."),
        ("Session?", "Flask session cookie stores user and last 20 plans."),
        ("CORS?", "flask-cors supports_credentials True."),
        ("Health?", "GET /health → mode demo|live."),
        ("Test count?", "22 passed, 83% coverage."),
        ("FR-04 mapping?", "TC-02,03,04,17,18."),
        ("Jewelry multimodal?", "Optional outfit bytes passed to generate_content parts."),
        ("Security of keys?", ".env gitignored; .env.example empty."),
        ("Why mock catalogs?", "Spec forbids scraping live sites."),
        ("Party split?", "Categories catering, decoration, entertainment, stay."),
        ("Home quantities?", "lights, fans, dining_tables, sofa, bed, storage."),
        ("Limitation?", "Prices are sample; videos shorter than brief."),
        ("Improve?", "Real APIs, DB, GST, PDF."),
        ("Where are prompts?", "services/recommend.py HOME_PROMPTS, PARTY_PROMPT, JEWEL_PROMPT."),
        ("Fallback message?", "ai_client.FALLBACK_MESSAGE shown in UI alert.warn."),
        ("Who is the user?", "Indian students/households planning events in rupees."),
    ]
    for i, (q, a) in enumerate(qa, 1):
        v.add_heading(f"Q{i}. {q}", 2)
        v.add_paragraph(a)
    save(v, PH / "08_Project_Demonstration" / "Viva_Questions.docx")


def demo_script():
    (ROOT / "docs" / "demo_script.md").write_text(
        """# PocketSmart AI narration

1. Landing (10s): Welcome to PocketSmart AI, a budget planner for home, parties and jewelry.
2. Register (15s): We create a sample account so history is stored in the session.
3. Home (40s): Budget eighty thousand rupees, living room, kitchen and bedroom with lights, fans and a dining table.
4. Home results (30s): Every card says sample data. The used total stays under budget.
5. Party (40s): Birthday, forty guests, fifty thousand rupees.
6. Jewelry (40s): Wedding guest, fifteen thousand, plus a generated outfit image.
7. History (20s): Session log of the three plans.
8. Edge (20s): Budget one hundred rupees is rejected.
""",
        encoding="utf-8",
    )


if __name__ == "__main__":
    phase1()
    phase2()
    phase3()
    phase4()
    phase5()
    phase6()
    phase7()
    phase8()
    demo_script()
    print("docs ok")
